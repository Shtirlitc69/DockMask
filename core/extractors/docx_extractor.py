"""DOCX экстрактор, извлекающий
текстовые блоки из DOCX документов."""

from pathlib import Path

from docx import Document

from core.context import enrich_context
from core.models import (
    BlockKind,
    DocumentFormat,
    ExtractedDocument,
    Location,
    TextBlock,
)


class DocxExtractor:
    """
    Задача - извлекать текстовые блоки из DOCX документов и
    представлять их в виде объектов TextBlock.
    """

    def extract(self, file_path: str | Path) -> ExtractedDocument:
        path = Path(file_path)
        document = Document(str(path))

        blocks: list[TextBlock] = []

        # Абзацы
        for paragraph_index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text
            if not text.strip():
                continue
            blocks.append(
                TextBlock(
                    block_id=f"docx_paragraph_{paragraph_index}",
                    text=text,
                    kind=BlockKind.DOCX_PARAGRAPH,
                    location=Location(paragraph_index=paragraph_index),
                )
            )

        # Таблицы
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    for cell_paragraph_index, paragraph in enumerate(cell.paragraphs):
                        text = paragraph.text
                        if not text.strip():
                            continue
                        blocks.append(
                            TextBlock(
                                block_id=(
                                    f"docx_table_{table_index}_r{row_index}_c{column_index}"
                                    f"_p{cell_paragraph_index}"
                                ),
                                text=text,
                                kind=BlockKind.DOCX_TABLE_CELL,
                                location=Location(
                                    table_index=table_index,
                                    row=row_index,
                                    column=column_index,
                                    cell_paragraph_index=cell_paragraph_index,
                                ),
                            )
                        )

        # Keep stable paragraph/table identifiers but restore OOXML document order.
        body_order = {node: index for index, node in enumerate(document.element.body)}
        paragraph_order = {i: body_order[p._p] for i, p in enumerate(document.paragraphs)}
        table_order = {i: body_order[t._tbl] for i, t in enumerate(document.tables)}
        blocks.sort(key=lambda b: (
            paragraph_order[b.location.paragraph_index]
            if b.location.paragraph_index is not None else table_order[b.location.table_index],
            b.location.row or 0, b.location.column or 0, b.location.cell_paragraph_index or 0,
        ))
        enrich_context(blocks)
        return ExtractedDocument(
            format=DocumentFormat.DOCX,
            blocks=blocks,
        )


def extract_docx(file_path: str | Path) -> ExtractedDocument:
    return DocxExtractor().extract(file_path)
