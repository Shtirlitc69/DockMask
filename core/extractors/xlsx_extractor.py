"""XLSX экстрактор для извлечения текстовых значений из ячеек XLSX файлов."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from core.context import enrich_context
from core.models import (
    BlockKind,
    DocumentFormat,
    ExtractedDocument,
    Location,
    TextBlock,
)


class XlsxExtractor:
    """Извлекает текстовые значения из ячеек XLSX."""

    def extract(self, file_path: str | Path) -> ExtractedDocument:
        path = Path(file_path)
        workbook = load_workbook(path, read_only=True, data_only=False)
        blocks: list[TextBlock] = []

        try:
            for sheet in workbook.worksheets:
                headers = {}
                for row_index, row in enumerate(sheet.iter_rows()):
                    for column_index, cell in enumerate(row):
                        value = cell.value
                        if value is None or cell.data_type == "f":
                            continue

                        text = self._normalize_cell_value(value)
                        if not text.strip():
                            continue

                        if row_index == 0:
                            headers[column_index] = text
                        blocks.append(
                            TextBlock(
                                block_id=f"xlsx_{sheet.title}_r{row_index}_c{column_index}_{cell.coordinate}",
                                text=text,
                                kind=BlockKind.XLSX_CELL,
                                context={"column_header": headers.get(column_index, ""), "row": row_index, "column": column_index},
                                location=Location(
                                    sheet_name=sheet.title,
                                    cell_coordinate=cell.coordinate,
                                ),
                            )
                        )
        finally:
            workbook.close()

        enrich_context(blocks)
        return ExtractedDocument(
            format=DocumentFormat.XLSX,
            blocks=blocks,
        )

    @staticmethod
    def _normalize_cell_value(value: Any) -> str:
        if value is None:
            return ""
        return str(value)


def extract_xlsx(file_path: str | Path) -> ExtractedDocument:
    return XlsxExtractor().extract(file_path)
