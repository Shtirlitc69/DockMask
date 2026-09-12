"""Build a safe, simplified preview from a completed redacted document."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

MAX_PREVIEW_CHARS = 500_000


class _Budget:
    def __init__(self, limit: int) -> None:
        self.remaining = limit
        self.truncated = False

    def take(self, text: str) -> str:
        if self.remaining <= 0:
            self.truncated = True
            return ""
        if len(text) <= self.remaining:
            self.remaining -= len(text)
            return text
        value = text[: self.remaining]
        self.remaining = 0
        self.truncated = True
        return value


def _segments(text: str, markers: tuple[str, ...]) -> list[dict[str, str | None]]:
    if not text:
        return [{"text": "", "replacement": None}]
    if not markers:
        return [{"text": text, "replacement": None}]
    pattern = re.compile("(" + "|".join(re.escape(item) for item in markers) + ")")
    result: list[dict[str, str | None]] = []
    for part in pattern.split(text):
        if part:
            result.append(
                {"text": part, "replacement": part if part in markers else None}
            )
    return result


def _line(
    element_id: str,
    text: str,
    markers: tuple[str, ...],
    budget: _Budget,
    *,
    line_type: str = "field",
) -> dict[str, Any] | None:
    value = budget.take(text)
    if not value and budget.truncated:
        return None
    return {"id": element_id, "type": line_type, "segments": _segments(value, markers)}


def _docx_preview(path: Path, markers: tuple[str, ...], budget: _Budget) -> list[dict[str, Any]]:
    document = Document(str(path))
    elements: list[dict[str, Any]] = []
    for index, item in enumerate(document.iter_inner_content()):
        if budget.truncated:
            break
        if isinstance(item, Paragraph):
            if not item.text.strip():
                continue
            style_name = item.style.name.casefold() if item.style is not None else ""
            line_type = "title" if style_name.startswith(("heading", "заголовок")) else "field"
            value = _line(f"docx-line-{index}", item.text, markers, budget, line_type=line_type)
            if value:
                elements.append(value)
            continue
        if isinstance(item, Table):
            rows: list[list[dict[str, Any]]] = []
            for row in item.rows:
                preview_row: list[dict[str, Any]] = []
                for cell in row.cells:
                    text = budget.take("\n".join(p.text for p in cell.paragraphs))
                    preview_row.append({"segments": _segments(text, markers)})
                    if budget.truncated:
                        break
                rows.append(preview_row)
                if budget.truncated:
                    break
            elements.append({"id": f"docx-table-{index}", "type": "table", "rows": rows})
    return elements


def _xlsx_preview(path: Path, markers: tuple[str, ...], budget: _Budget) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    elements: list[dict[str, Any]] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            heading = _line(
                f"xlsx-sheet-{sheet_index}",
                f"Лист «{sheet.title}»",
                markers,
                budget,
                line_type="title",
            )
            if heading:
                elements.append(heading)
            rows: list[list[dict[str, Any]]] = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if not values:
                    continue
                preview_row: list[dict[str, Any]] = []
                for value in values:
                    text = budget.take(value)
                    preview_row.append({"segments": _segments(text, markers)})
                    if budget.truncated:
                        break
                rows.append(preview_row)
                if budget.truncated:
                    break
            if rows:
                elements.append(
                    {"id": f"xlsx-table-{sheet_index}", "type": "table", "rows": rows}
                )
            if budget.truncated:
                break
    finally:
        workbook.close()
    return elements


def _pdf_preview(path: Path, markers: tuple[str, ...], budget: _Budget) -> list[dict[str, Any]]:
    elements: list[dict[str, Any]] = []
    with fitz.open(path) as document:
        for page_index, page in enumerate(document):
            heading = _line(
                f"pdf-page-{page_index}",
                f"Страница {page_index + 1}",
                markers,
                budget,
                line_type="title",
            )
            if heading:
                elements.append(heading)
            for block_index, raw in enumerate(page.get_text("blocks", sort=True)):
                text = raw[4] if len(raw) > 4 else ""
                if not isinstance(text, str) or not text.strip():
                    continue
                value = _line(
                    f"pdf-page-{page_index}-block-{block_index}",
                    text.strip(),
                    markers,
                    budget,
                )
                if value:
                    elements.append(value)
                if budget.truncated:
                    break
            if budget.truncated:
                break
    return elements


def build_preview(
    path: str | Path,
    replacements: list[str],
    *,
    limit: int = MAX_PREVIEW_CHARS,
    marker_aliases: dict[str, str] | None = None,
) -> tuple[list[dict[str, Any]], bool]:
    source = Path(path)
    aliases = marker_aliases if marker_aliases is not None else {item: item for item in replacements}
    markers = tuple(sorted((item for item in aliases if item), key=len, reverse=True))
    budget = _Budget(limit)
    suffix = source.suffix.casefold()
    if suffix == ".docx":
        elements = _docx_preview(source, markers, budget)
    elif suffix == ".xlsx":
        elements = _xlsx_preview(source, markers, budget)
    elif suffix == ".pdf":
        elements = _pdf_preview(source, markers, budget)
    else:
        raise ValueError("unsupported preview format")
    for element in elements:
        groups = [element.get("segments", [])]
        groups.extend(cell["segments"] for row in element.get("rows", []) for cell in row)
        for group in groups:
            for segment in group:
                if segment.get("replacement"):
                    segment["replacement"] = aliases[segment["replacement"]]
    return elements, budget.truncated
