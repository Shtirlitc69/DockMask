"""XLSX redaction with cell-style preservation and visible highlighting."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from core.models import Match

HIGHLIGHT_FILL = PatternFill("solid", fgColor="FFF59D")


def redact_xlsx(
    source_file: str | Path,
    output_file: str | Path,
    matches_by_block: dict[str, list[Match]],
) -> Path:
    source = Path(source_file)
    target = Path(output_file)
    if source.resolve() == target.resolve():
        raise ValueError("output_file must differ from source_file")
    workbook = load_workbook(source, data_only=False)
    try:
        for matches in matches_by_block.values():
            for match in sorted(matches, key=lambda item: (item.start, item.end), reverse=True):
                location = match.location
                if not location.sheet_name or not location.cell_coordinate:
                    continue
                if location.sheet_name not in workbook.sheetnames:
                    continue
                cell = workbook[location.sheet_name][location.cell_coordinate]
                if cell.data_type == "f" or cell.value is None:
                    continue
                text = str(cell.value)
                if not 0 <= match.start < match.end <= len(text):
                    continue
                if text[match.start : match.end] != match.text:
                    continue
                cell.value = text[: match.start] + match.replacement + text[match.end :]
                cell.fill = HIGHLIGHT_FILL
                match.applied = True
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target)
    finally:
        workbook.close()
    return target
