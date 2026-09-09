"""DOCX redaction that preserves unaffected runs and document structure."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from core.models import Match


def _clone_run_after(paragraph: Paragraph, anchor: Run, text: str) -> Run:
    clone = paragraph.add_run(text)
    if anchor._r.rPr is not None:
        clone._r.insert(0, deepcopy(anchor._r.rPr))
    anchor._r.addnext(clone._r)
    return clone


def _run_position(runs: list[Run], offset: int, *, end: bool = False) -> tuple[int, int]:
    cursor = 0
    for index, run in enumerate(runs):
        next_cursor = cursor + len(run.text)
        if offset < next_cursor or (end and offset == next_cursor):
            return index, offset - cursor
        cursor = next_cursor
    raise ValueError("span offset is outside paragraph runs")


class DocxRedactor:
    """Apply exact redactions without overwriting the source file."""

    def redact(
        self,
        source_file: str | Path,
        output_file: str | Path,
        matches_by_block: dict[str, list[Match]],
    ) -> Path:
        source_path = Path(source_file)
        target_path = Path(output_file)
        if source_path.resolve() == target_path.resolve():
            raise ValueError("output_file must differ from source_file")
        document = Document(str(source_path))

        for block_id, matches in matches_by_block.items():
            paragraph = self._resolve_paragraph(document, block_id)
            if paragraph is not None:
                self._apply_to_paragraph(paragraph, matches)

        target_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(target_path))
        return target_path

    @staticmethod
    def _resolve_paragraph(document: Document, block_id: str) -> Paragraph | None:
        if block_id.startswith("docx_paragraph_"):
            try:
                index = int(block_id.removeprefix("docx_paragraph_"))
                return document.paragraphs[index]
            except (ValueError, IndexError):
                return None

        if not block_id.startswith("docx_table_"):
            return None
        try:
            table_part, paragraph_part = block_id.removeprefix("docx_table_").rsplit("_p", 1)
            table_text, rest = table_part.split("_r", 1)
            row_text, column_text = rest.split("_c", 1)
            cell = document.tables[int(table_text)].rows[int(row_text)].cells[int(column_text)]
            return cell.paragraphs[int(paragraph_part)]
        except (ValueError, IndexError):
            return None

    @staticmethod
    def _apply_to_paragraph(paragraph: Paragraph, matches: list[Match]) -> None:
        accepted: list[Match] = []
        occupied: list[tuple[int, int]] = []
        source_text = paragraph.text
        for match in sorted(matches, key=lambda item: (item.start, item.end), reverse=True):
            if not 0 <= match.start < match.end <= len(source_text):
                continue
            if source_text[match.start : match.end] != match.text:
                continue
            if any(match.start < end and start < match.end for start, end in occupied):
                continue
            if DocxRedactor._replace_span(paragraph, match):
                occupied.append((match.start, match.end))
                accepted.append(match)
        for match in accepted:
            match.applied = True

    @staticmethod
    def _replace_span(paragraph: Paragraph, match: Match) -> bool:
        runs = list(paragraph.runs)
        if not runs:
            return False
        try:
            start_index, start_offset = _run_position(runs, match.start)
            end_index, end_offset = _run_position(runs, match.end, end=True)
        except ValueError:
            return False

        start_run = runs[start_index]
        end_run = runs[end_index]
        prefix = start_run.text[:start_offset]
        suffix = end_run.text[end_offset:]
        start_run.text = prefix

        if start_index == end_index:
            suffix_run = _clone_run_after(paragraph, start_run, suffix)
            replacement_run = _clone_run_after(paragraph, start_run, match.replacement)
            replacement_run._r.addnext(suffix_run._r)
        else:
            for run in runs[start_index + 1 : end_index]:
                run.text = ""
            end_run.text = suffix
            replacement_run = _clone_run_after(paragraph, start_run, match.replacement)

        replacement_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        replacement_run.font.bold = True
        return True


def redact_docx(
    source_file: str | Path,
    output_file: str | Path,
    matches_by_block: dict[str, list[Match]],
) -> Path:
    return DocxRedactor().redact(source_file, output_file, matches_by_block)
