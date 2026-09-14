import json
import re
from pathlib import Path

import fitz
import pytest
from docx import Document

from core.detectors.entity_detector import detect_all
from core.detectors.party_identifier import identify_parties
from core.detectors.rule_based import detect_rule_based
from core.extractors.docx_extractor import extract_docx
from core.extractors.pdf_extractor import extract_pdf
from core.models import EntityType, JobStatus, PartyRole
from core.orchestrator import run_pipeline
from llm.mock_client import MockLLMClient

FIXTURES = Path(__file__).parent / "fixtures" / "masking"


class EmptyClient(MockLLMClient):
    async def find_entities(self, text, types):
        return []

    async def classify_party(self, context_snippet, candidate_name):
        return None


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "text,kind,expected,needs_review",
    [
        ("/ Смирнов А.В. /", EntityType.PERSON_NAME, ["Смирнов А.В."], False),
        ("г. Москва, 12 марта 2026 г.", EntityType.ADDRESS, ["г. Москва, 12 марта 2026 г."], True),
    ],
)
async def test_model_boundaries_and_semantic_review(text, kind, expected, needs_review):
    from core.models import BlockKind, EntitySpan, Location, TextBlock

    class ExactClient(EmptyClient):
        async def find_entities(self, source, types):
            return [EntitySpan(kind, source, 0, len(source), confidence=0.8)]

    block = TextBlock("sample", text, BlockKind.DOCX_PARAGRAPH, Location(paragraph_index=0))
    matches = (await detect_all([block], [kind], ExactClient()))["sample"]
    assert [match.text for match in matches] == expected
    assert all(bool(match.review_reason) == needs_review for match in matches)


def pattern(value):
    return re.compile(r"\s+".join(map(re.escape, value.split())))


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "filename,extract", [("contract.docx", extract_docx), ("invoice.pdf", extract_pdf)]
)
async def test_synthetic_document_has_no_known_leaks_without_model(filename, extract, tmp_path):
    blocks = extract(FIXTURES / filename).blocks
    matches = await detect_all(blocks, list(EntityType), EmptyClient())
    expected = json.loads((FIXTURES / "expected.json").read_text(encoding="utf-8"))[filename]
    for kind, values in expected.items():
        for value in values:
            occurrences = [
                (block, occurrence)
                for block in blocks
                for occurrence in pattern(value).finditer(block.text)
            ]
            assert occurrences, (kind, value)
            for block, occurrence in occurrences:
                assert any(
                    m.entity_type.value == kind
                    and m.start <= occurrence.start()
                    and m.end >= occurrence.end()
                    for m in matches[block.block_id]
                ), (block.block_id, kind, value)
    await identify_parties(blocks, matches, EmptyClient())
    for group in matches.values():
        for m in group:
            if "Промсвязьбанк" in m.text or m.text in {"Смирнов А.В.", "Фёдорова Ольга Петровна"}:
                assert m.party_role is PartyRole.SUPPLIER
    if filename == "contract.docx":
        flat = [m for group in matches.values() for m in group]
        person = next(m for m in flat if m.text == "Ковалёвой Ирины Николаевны")
        school = next(m for m in flat if m.text == "МБОУ СОШ № 42")
        assert person.organization_id == school.entity_id
        assert (
            next(m for m in flat if m.text.startswith("Общество с")).entity_id
            == next(m for m in flat if m.text == "ООО «ТехноКомпонент»").entity_id
        )

    events = []

    async def on_progress(*event):
        events.append(event)

    result = await run_pipeline(
        FIXTURES / filename,
        tmp_path,
        list(EntityType),
        EmptyClient(),
        progress_callback=on_progress,
    )
    assert result.status is JobStatus.DONE
    assert [e[0] for e in events if e[1] == "done"] == [
        "extract",
        "detect",
        "clarify",
        "redact",
        "report",
    ]
    if filename.endswith("pdf"):
        with fitz.open(result.output_document) as pdf:
            remaining = "\n".join(p.get_text() for p in pdf)
    else:
        remaining = "\n".join(b.text for b in extract_docx(result.output_document).blocks)
    for values in expected.values():
        for value in values:
            assert not pattern(value).search(remaining), value
    for preserved in (
        ["Генеральный директор", "20%", "0,1%"]
        if filename.endswith("docx")
        else ["5 000", "20%", "Руководитель", "Главный бухгалтер"]
    ):
        assert preserved in remaining
    diagnostic = json.loads(
        (tmp_path / (Path(filename).stem + ".diagnostics.json")).read_text(encoding="utf-8")
    )
    encoded = json.dumps(diagnostic, ensure_ascii=False)
    assert "checksum_valid" in encoded
    assert "7719123456" not in encoded
    assert "Петров" not in encoded


@pytest.mark.parametrize(
    "text",
    [
        "Генеральный директор",
        "Главный бухгалтер",
        "20%",
        "0,1%",
        "12.03.2026",
        "5 000 шт.",
        "№ 128",
    ],
)
def test_not_person_or_money(text):
    assert not detect_rule_based(text, [EntityType.PERSON_NAME, EntityType.AMOUNT])


def test_docx_order_and_column_context(tmp_path):
    doc = Document()
    doc.add_paragraph("До таблицы")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Поставщик"
    table.cell(0, 1).text = "Покупатель"
    table.cell(1, 0).text = "Иванов И.И."
    table.cell(1, 1).text = "Петров П.П."
    doc.add_paragraph("После таблицы")
    path = tmp_path / "order.docx"
    doc.save(path)
    blocks = extract_docx(path).blocks
    assert blocks[-1].text == "После таблицы"
    assert (
        next(b for b in blocks if b.text == "Иванов И.И.").context["column_header"] == "Поставщик"
    )


@pytest.mark.asyncio
async def test_unknown_role_does_not_block_obvious_mask(tmp_path):
    doc = Document()
    doc.add_paragraph("ООО Тест")
    source = tmp_path / "source.docx"
    doc.save(source)
    result = await run_pipeline(
        source, tmp_path / "result", [EntityType.ORGANIZATION], MockLLMClient()
    )
    assert result.status is JobStatus.DONE
    assert not result.open_questions


def test_money_words_do_not_swallow_surrounding_sentence():
    spans = detect_rule_based("Два товара стоят сто рублей", [EntityType.AMOUNT])
    assert [s.text for s in spans] == ["сто рублей"]


def test_grouped_identifiers_and_nbsp_preserve_offsets():
    text = "ИНН\n77\u00a019123456; ОГРН 114\u202f7746123456"
    spans = detect_rule_based(text, [EntityType.INN, EntityType.OGRN])
    assert len(spans) == 2
    assert all(text[s.start : s.end] == s.text for s in spans)


@pytest.mark.asyncio
async def test_phone_cannot_be_assembled_from_independent_table_cells():
    from core.models import BlockKind, Location, TextBlock

    text = "тел. +7 (916)\n555-01-23"
    block = TextBlock(
        "pdf",
        text,
        BlockKind.PDF_TEXT_BLOCK,
        Location(page_number=0),
        context={
            "cell_spans": [
                {"start": 0, "end": 13, "table": 0, "row": 1, "column": 0},
                {"start": 14, "end": len(text), "table": 0, "row": 1, "column": 1},
            ],
        },
    )
    result = await detect_all([block], [EntityType.PHONE], EmptyClient())
    assert result["pdf"] == []


@pytest.mark.asyncio
async def test_money_table_value_does_not_mask_quantity_elsewhere():
    from core.models import BlockKind, Location, TextBlock

    blocks = [
        TextBlock(
            "price",
            "100",
            BlockKind.XLSX_CELL,
            Location(sheet_name="Sheet", cell_coordinate="B2"),
            context={"column_header": "Цена"},
        ),
        TextBlock(
            "quantity",
            "100",
            BlockKind.XLSX_CELL,
            Location(sheet_name="Sheet", cell_coordinate="A2"),
            context={"column_header": "Количество"},
        ),
    ]
    result = await detect_all(blocks, [EntityType.AMOUNT], EmptyClient())
    assert len(result["price"]) == 1
    assert result["quantity"] == []
