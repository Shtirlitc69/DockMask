"""Regressions for multiline PDF deletion, label geometry and missed repeats."""

from pathlib import Path

import fitz
import pytest

from core.detectors.entity_detector import detect_all
from core.extractors.pdf_extractor import extract_pdf
from core.models import BlockKind, EntitySpan, EntityType, Location, Match, TextBlock
from core.redaction.pdf_redactor import redact_pdf
from llm.mock_client import MockLLMClient
from llm.structured import EntityItem, to_entity_spans


@pytest.mark.parametrize(
    "style,label", [("full", "[ОРГАНИЗАЦИЯ_1]"), ("short", "[ОРГ_1]"), ("none", "")]
)
def test_multiline_pdf_deletes_all_lines_and_centres_label(tmp_path: Path, style, label):
    source, target = tmp_path / "source.pdf", tmp_path / "result.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((50, 60), "Before Long Company\nName After", fontsize=12)
        doc.save(source)
    block = extract_pdf(source).blocks[0]
    text = "Long Company\nName"
    match = Match(
        block.block_id,
        EntityType.ORGANIZATION,
        text,
        block.text.index(text),
        block.text.index(text) + len(text),
        "[ОРГАНИЗАЦИЯ_1]",
        "test",
        1,
        block.location,
    )
    redact_pdf(source, target, {block.block_id: [match]}, label_style=style)
    with fitz.open(target) as doc:
        page = doc[0]
        assert "Company" not in page.get_text()
        assert "Name" not in page.get_text()
        assert "Before" in page.get_text() and "After" in page.get_text()
        assert match.applied
        if label:
            assert label in page.get_text()
            drawn = page.search_for(label)[0]
            with fitz.open(source) as original:
                rect = original[0].search_for("Long Company")[0]
            assert abs((drawn.x0 + drawn.x1) - (rect.x0 + rect.x1)) < 1
            assert abs((drawn.y0 + drawn.y1) - (rect.y0 + rect.y1)) < 1
        else:
            assert "[" not in page.get_text()


def test_structured_response_recovers_linewrap_without_matching_longer_words():
    text = "Иванов Иван\nИванович. Иванов Иван Ивановичев."
    spans = to_entity_spans(
        text,
        frozenset([EntityType.PERSON_NAME]),
        [EntityItem(type="person_name", text="Иванов Иван Иванович")],
        source="test",
    )
    assert len(spans) == 1
    assert spans[0].text == "Иванов Иван\nИванович"
    assert text[spans[0].start : spans[0].end] == spans[0].text


@pytest.mark.asyncio
async def test_document_repeats_across_batches_keep_marker_and_word_boundaries():
    class MissingRepeatsClient(MockLLMClient):
        async def find_entities(self, text, types):
            if not text.startswith("First"):
                return []
            return [EntitySpan(EntityType.PERSON_NAME, "Alice Example", 6, 19)]

    blocks = [
        TextBlock(str(i), text, BlockKind.PDF_TEXT_BLOCK, Location(page_number=i))
        for i, text in enumerate(
            ["First Alice Example", "Alice\nExample", "Alice Exampleton", "Alice Example"]
        )
    ]
    result = await detect_all(blocks, [EntityType.PERSON_NAME], MissingRepeatsClient())
    assert [len(result[b.block_id]) for b in blocks] == [1, 1, 0, 1]
    assert len({m.replacement for group in result.values() for m in group}) == 1
    assert result["1"][0].source == "document_repeat"


@pytest.mark.parametrize("style", ["short", "none"])
def test_label_choice_survives_job_storage_and_worker(tmp_path, style):
    import time
    from io import BytesIO

    from docx import Document
    from fastapi.testclient import TestClient

    from app.main import create_app
    from core.config import Settings
    from tests.test_runtime_e2e import MemorySecretStore

    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("До contact@example.test после")
    doc.save(source)
    app = create_app(
        settings=Settings(data_dir=tmp_path / "runtime"), secret_store=MemorySecretStore()
    )
    with TestClient(app) as client:
        response = client.post(
            "/api/jobs",
            files={"file": (source.name, source.read_bytes())},
            data={"entity_types": "email", "label_style": style},
        )
        assert response.status_code == 200
        job = response.json()["job_id"]
        for _ in range(100):
            state = client.get(f"/api/jobs/{job}").json()["status"]
            if state in {"done", "failed"}:
                break
            time.sleep(0.03)
        assert state == "done"
        report = client.get(f"/api/jobs/{job}/report").json()
        assert report["replacements"][0]["replacement"] == "[ЭЛЕКТРОННАЯ_ПОЧТА_1]"
        response = client.get(f"/api/jobs/{job}/document")
        assert response.status_code == 200
        text = Document(BytesIO(response.content)).paragraphs[0].text
        assert "contact@example.test" not in text
        assert "До " in text and " после" in text
        assert ("[ПОЧТА_1]" in text) if style == "short" else ("[" not in text)
        preview = client.get(f"/api/jobs/{job}/preview").json()
        segments = [s for e in preview["elements"] for s in e.get("segments", [])]
        labelled = [s for s in segments if s.get("replacement")]
        if style == "short":
            assert labelled == [{"text": "[ПОЧТА_1]", "replacement": "[ЭЛЕКТРОННАЯ_ПОЧТА_1]"}]
        else:
            assert not labelled
