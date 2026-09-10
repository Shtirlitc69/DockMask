import asyncio
import threading
import time
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.main import create_app
from core.config import Settings


class MemorySecretStore:
    def __init__(self) -> None:
        self.values: dict[str, str] = {}

    async def get(self, provider: str) -> str | None:
        return self.values.get(provider)

    async def set(self, provider: str, value: str) -> None:
        self.values[provider] = value

    async def delete(self, provider: str) -> None:
        self.values.pop(provider, None)


def test_runtime_cancel_interrupts_active_pipeline(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("ООО Тест")
    document.save(source)
    started = threading.Event()
    cancelled = threading.Event()

    async def slow_pipeline(*_args: object, **_kwargs: object) -> object:
        started.set()
        try:
            await asyncio.Event().wait()
        except asyncio.CancelledError:
            cancelled.set()
            raise

    app = create_app(
        settings=Settings(data_dir=tmp_path / "runtime", env="development"),
        secret_store=MemorySecretStore(),
    )
    with (
        patch("app.runtime.run_pipeline", new=slow_pipeline),
        TestClient(app) as client,
        source.open("rb") as stream,
    ):
        created = client.post(
            "/api/jobs",
            files={"file": (source.name, stream)},
            data={"entity_types": "organization"},
        )
        job_id = created.json()["job_id"]
        assert started.wait(timeout=2)

        response = client.post(f"/api/jobs/{job_id}/cancel")

        assert response.status_code == 200
        assert response.json()["status"] == "cancelled"
        assert cancelled.wait(timeout=2)


def test_production_runtime_processes_docx_with_default_local_provider(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("E-mail: email@example.test")
    document.save(source)
    secret_store = MemorySecretStore()
    app = create_app(
        settings=Settings(data_dir=tmp_path / "runtime", env="production"),
        secret_store=secret_store,
    )

    with TestClient(app) as client, source.open("rb") as stream:
        config = client.get("/api/config")
        assert config.status_code == 200
        payload = config.json()
        assert payload["provider"] == "mock"
        assert payload["model"] == "mock"
        local_provider = next(item for item in payload["providers"] if item["id"] == "mock")
        assert local_provider == {
            "id": "mock",
            "display_name": "Локальный режим (без LLM)",
            "available": True,
            "requires_base_url": False,
            "api_key_optional": True,
            "development_only": False,
            "recommended_models": ["mock"],
        }
        created = client.post(
            "/api/jobs",
            files={"file": (source.name, stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"entity_types": "email"},
        )
        assert created.status_code == 200
        job_id = created.json()["job_id"]
        status = None
        for _ in range(100):
            status = client.get(f"/api/jobs/{job_id}").json()
            if status["status"] in {"done", "failed", "needs_clarification"}:
                break
            time.sleep(0.02)

        assert status is not None and status["status"] == "done"
        assert status["total_replacements"] == 1
        report = client.get(f"/api/jobs/{job_id}/report")
        json_report = client.get(f"/api/jobs/{job_id}/report.json")
        csv_report = client.get(f"/api/jobs/{job_id}/report.csv")
        preview = client.get(f"/api/jobs/{job_id}/preview")
        xlsx = client.get(f"/api/jobs/{job_id}/report.xlsx")
        result = client.get(f"/api/jobs/{job_id}/document")
        assert (
            report.status_code
            == json_report.status_code
            == csv_report.status_code
            == preview.status_code
            == xlsx.status_code
            == result.status_code
            == 200
        )
        assert report.json()["total_replacements"] == 1
        assert preview.json()["elements"]
        assert "[ЭЛЕКТРОННАЯ_ПОЧТА_1]" in json_report.text
        assert "Электронная почта" in csv_report.content.decode("utf-8-sig")
        assert b"email@example.test" not in result.content


def test_provider_keys_are_isolated_and_write_only(tmp_path: Path) -> None:
    secret_store = MemorySecretStore()
    app = create_app(
        settings=Settings(data_dir=tmp_path / "runtime", env="development"),
        secret_store=secret_store,
    )
    first = "openai-secret-value"
    second = "anthropic-secret-value"

    with TestClient(app) as client:
        assert client.put(
            "/api/config",
            json={"provider": "openai", "model": "test", "api_key": first},
        ).status_code == 200
        response = client.put(
            "/api/config",
            json={"provider": "anthropic", "model": "test", "api_key": second},
        )
        fetched = client.get("/api/config")

    assert response.status_code == 200
    assert first not in response.text and second not in response.text
    assert first not in fetched.text and second not in fetched.text
    assert secret_store.values == {"openai": first, "anthropic": second}


def test_runtime_cancels_clarification_and_removes_job_files(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("ООО Тест")
    document.save(source)
    data_dir = tmp_path / "runtime"
    app = create_app(
        settings=Settings(data_dir=data_dir, env="development"),
        secret_store=MemorySecretStore(),
    )

    with TestClient(app) as client, source.open("rb") as stream:
        created = client.post(
            "/api/jobs",
            files={"file": (source.name, stream)},
            data={"entity_types": "organization"},
        )
        job_id = created.json()["job_id"]
        for _ in range(100):
            status = client.get(f"/api/jobs/{job_id}").json()
            if status["status"] == "needs_clarification":
                break
            time.sleep(0.02)

        cancelled = client.post(f"/api/jobs/{job_id}/cancel")

        assert status["status"] == "needs_clarification"
        assert cancelled.status_code == 200
        assert cancelled.json()["status"] == "cancelled"
        assert client.get(f"/api/jobs/{job_id}").json()["status"] == "cancelled"
        assert client.get(f"/api/jobs/{job_id}/document").status_code == 409
        assert not (data_dir / "jobs" / job_id).exists()
