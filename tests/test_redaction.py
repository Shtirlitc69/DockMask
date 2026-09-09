"""Redaction tests for DOCX, XLSX and text PDF."""

from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook

from core.models import EntityType, Location, Match
from core.redaction.docx_redactor import redact_docx
from core.redaction.pdf_redactor import redact_pdf
from core.redaction.xlsx_redactor import redact_xlsx


def _match(block_id: str, text: str, start: int, location: Location) -> Match:
    return Match(
        block_id=block_id,
        entity_type=EntityType.EMAIL,
        text=text,
        start=start,
        end=start + len(text),
        replacement="[EMAIL_1]",
        source="rule",
        confidence=1,
        location=location,
    )


def test_docx_replaces_span_across_runs_and_preserves_surroundings(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    target = tmp_path / "result.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("До ").bold = True
    paragraph.add_run("email@").italic = True
    paragraph.add_run("example.test после").underline = True
    document.save(source)
    match = _match("docx_paragraph_0", "email@example.test", 3, Location(paragraph_index=0))

    redact_docx(source, target, {match.block_id: [match]})

    result = Document(target)
    assert result.paragraphs[0].text == "До [EMAIL_1] после"
    assert result.paragraphs[0].runs[0].bold is True
    assert result.paragraphs[0].runs[-1].underline is True
    assert match.applied is True


def test_xlsx_keeps_formula_and_highlights_changed_cell(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    target = tmp_path / "result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Данные"
    sheet["A1"] = "До email@example.test после"
    sheet["B1"] = "=1+1"
    workbook.save(source)
    workbook.close()
    match = _match(
        "xlsx_Данные_r0_c0_A1",
        "email@example.test",
        3,
        Location(sheet_name="Данные", cell_coordinate="A1"),
    )

    redact_xlsx(source, target, {match.block_id: [match]})

    result = load_workbook(target, data_only=False)
    assert result["Данные"]["A1"].value == "До [EMAIL_1] после"
    assert result["Данные"]["A1"].fill.fill_type == "solid"
    assert result["Данные"]["B1"].value == "=1+1"
    result.close()
    assert match.applied is True


def test_pdf_removes_original_text_and_adds_replacement(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    target = tmp_path / "result.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "email@example.test")
    document.save(source)
    document.close()
    source_doc = fitz.open(source)
    block = source_doc[0].get_text("blocks")[0]
    source_doc.close()
    match = _match(
        "pdf_page_0_block_0",
        "email@example.test",
        0,
        Location(page_number=0, bbox=tuple(float(v) for v in block[:4])),
    )

    redact_pdf(source, target, {match.block_id: [match]})

    result = fitz.open(target)
    text = "".join(page.get_text() for page in result)
    result.close()
    assert "email@example.test" not in text
    assert "[EMAIL_1]" in text
    assert match.applied is True
