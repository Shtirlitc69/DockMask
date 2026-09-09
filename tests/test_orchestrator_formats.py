import json
from pathlib import Path

import fitz
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from core.models import EntityType, JobStatus
from core.orchestrator import run_pipeline
from llm.mock_client import MockLLMClient


class CountingMockClient(MockLLMClient):
    def __init__(self) -> None:
        self.entity_calls = 0
        self.party_calls = 0

    async def find_entities(self, text, types):
        self.entity_calls += 1
        return await super().find_entities(text, types)

    async def classify_party(self, context_snippet, candidate_name):
        self.party_calls += 1
        return await super().classify_party(context_snippet, candidate_name)


@pytest.mark.asyncio
async def test_docx_pipeline_outputs_document_and_two_reports(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Телефон: +7 999 123-45-67")
    document.save(source)

    result = await run_pipeline(source, tmp_path / "result", [EntityType.PHONE], MockLLMClient())

    assert result.status is JobStatus.DONE
    assert len(result.matches) == 1
    assert result.output_document and result.report and result.xlsx_report
    redacted = Document(result.output_document)
    assert "+7 999 123-45-67" not in redacted.paragraphs[0].text
    assert "[PHONE_1]" in redacted.paragraphs[0].text
    assert len(json.loads(Path(result.report).read_text(encoding="utf-8"))) == 1
    workbook = load_workbook(result.xlsx_report)
    assert workbook.active.max_row == 2
    workbook.close()


@pytest.mark.asyncio
async def test_xlsx_pipeline_keeps_formula_and_redacts_text(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "email@example.test"
    sheet["B1"] = "=1+1"
    workbook.save(source)
    workbook.close()

    result = await run_pipeline(source, tmp_path / "result", [EntityType.EMAIL], MockLLMClient())

    assert result.status is JobStatus.DONE
    redacted = load_workbook(result.output_document, data_only=False)
    assert redacted.active["A1"].value == "[EMAIL_1]"
    assert redacted.active["B1"].value == "=1+1"
    redacted.close()


@pytest.mark.asyncio
async def test_text_pdf_pipeline_performs_real_redaction(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "email@example.test")
    document.save(source)
    document.close()

    result = await run_pipeline(source, tmp_path / "result", [EntityType.EMAIL], MockLLMClient())

    assert result.status is JobStatus.DONE
    redacted = fitz.open(result.output_document)
    text = "".join(page.get_text() for page in redacted)
    redacted.close()
    assert "email@example.test" not in text
    assert "[EMAIL_1]" in text


@pytest.mark.asyncio
async def test_scanned_pdf_reports_ocr_as_optional_unavailable_feature(tmp_path: Path) -> None:
    source = tmp_path / "scan.pdf"
    document = fitz.open()
    document.new_page()
    document.save(source)
    document.close()

    without_ocr = await run_pipeline(source, tmp_path / "a", [EntityType.EMAIL], MockLLMClient())
    with_ocr = await run_pipeline(
        source,
        tmp_path / "b",
        [EntityType.EMAIL],
        MockLLMClient(),
        use_ocr=True,
    )

    assert without_ocr.error_message == "ocr_required"
    assert with_ocr.error_message == "ocr_not_available"


@pytest.mark.asyncio
async def test_clarification_resumes_without_second_llm_call(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("ООО Тест")
    document.save(source)
    client = CountingMockClient()

    first = await run_pipeline(source, tmp_path / "result", [EntityType.ORGANIZATION], client)
    assert first.status is JobStatus.NEEDS_CLARIFICATION
    calls = (client.entity_calls, client.party_calls)
    question = first.open_questions[0]
    second = await run_pipeline(
        source,
        tmp_path / "result",
        [EntityType.ORGANIZATION],
        client,
        answers={question.question_id: "unknown"},
        prepared_matches=first.matches,
    )

    assert second.status is JobStatus.DONE
    assert (client.entity_calls, client.party_calls) == calls
