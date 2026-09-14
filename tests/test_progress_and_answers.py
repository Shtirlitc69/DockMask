import asyncio

import pytest
from docx import Document

from app.runtime import _deserialize_matches, _serialize_matches
from core.models import DocumentFormat, EntityType, JobStatus, PartyRole
from core.orchestrator import run_pipeline
from llm.types import LLMClientConfig, ProviderId
from storage.db import Database
from storage.jobs_repository import JobsRepository
from tests.test_orchestrator_formats import UncertainClient


async def repository(tmp_path):
    db = Database(tmp_path / "jobs.sqlite")
    await db.connect()
    repo = JobsRepository(db)
    await repo.create_job(
        job_id="abc",
        source_filename="source.docx",
        document_format=DocumentFormat.DOCX,
        entity_types=(EntityType.PERSON_NAME,),
        input_path=str(tmp_path / "source.docx"),
        config=LLMClientConfig(provider=ProviderId.MOCK, model="mock"),
    )
    await repo.set_processing("abc")
    return db, repo


@pytest.mark.asyncio
async def test_real_progress_wait_resume_persistence_and_keep(tmp_path):
    db, repo = await repository(tmp_path)
    document = Document()
    document.add_paragraph("Контакт: Иван Петров")
    document.add_paragraph("Иван Петров")
    document.save(tmp_path / "source.docx")
    started, release = asyncio.Event(), asyncio.Event()

    class SlowClient(UncertainClient):
        async def find_entities(self, text, types):
            started.set()
            await release.wait()
            return await super().find_entities(text, types)

    client = SlowClient()

    async def callback(*event):
        await repo.update_stage("abc", *event)

    task = asyncio.create_task(
        run_pipeline(
            tmp_path / "source.docx",
            tmp_path / "result",
            [EntityType.PERSON_NAME],
            client,
            progress_callback=callback,
        )
    )
    try:
        await asyncio.wait_for(started.wait(), 3)
        running = await repo.get_job("abc")
        assert running.stages[0]["status"] == "done"
        assert running.stages[1]["status"] == "active"
        assert running.progress == 20
        release.set()
        first = await task
        assert first.status is JobStatus.NEEDS_CLARIFICATION
        assert len(first.open_questions) == 1
        question = first.open_questions[0]
        assert question.kind == "mask_decision"
        assert len(question.contexts) == 2
        pending = _serialize_matches(first.matches)
        assert pending[0]["review_reason"]
        await repo.set_questions("abc", [{"question_id": question.question_id}], pending)
        waiting = await repo.get_job("abc")
        assert waiting.stages[2]["status"] == "waiting"
        assert waiting.progress == 40
        times = [s["finished_at"] for s in waiting.stages[:2]]
        await db.close()
        await db.connect()
        assert (await repo.get_job("abc")).stages == waiting.stages
        answers = {question.question_id: "keep"}
        await repo.save_answers("abc", answers)
        await repo.set_processing("abc")
        calls = client.entity_calls
        second = await run_pipeline(
            tmp_path / "source.docx",
            tmp_path / "result",
            [EntityType.PERSON_NAME],
            client,
            answers=answers,
            prepared_matches=_deserialize_matches(pending),
            progress_callback=callback,
        )
        assert second.status is JobStatus.DONE
        assert second.matches == []
        assert client.entity_calls == calls
        assert [s["finished_at"] for s in (await repo.get_job("abc")).stages[:2]] == times
        assert "Иван Петров" in Document(second.output_document).paragraphs[0].text
    finally:
        release.set()
        await asyncio.gather(task, return_exceptions=True)
        await db.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("terminal", ["failed", "cancelled"])
async def test_failure_and_cancel_never_complete_steps(tmp_path, terminal):
    db, repo = await repository(tmp_path)
    try:
        await repo.update_stage("abc", "extract", "done")
        await repo.update_stage("abc", "detect", "active", "Ожидание модели")
        if terminal == "failed":
            await repo.set_failed("abc", "provider_unavailable")
        else:
            await repo.set_cancelled("abc")
        await repo.update_stage(
            "abc", "detect", "done"
        )  # Late callback cannot undo terminal state.
        result = await repo.get_job("abc")
        assert result.progress == 20
        assert result.stages[1]["status"] == ("error" if terminal == "failed" else "cancelled")
        assert result.stages[2]["status"] == "pending"
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_old_role_answer_is_still_accepted(tmp_path):
    from core.models import Location, Match

    doc = Document()
    doc.add_paragraph("ООО Тест")
    path = tmp_path / "source.docx"
    doc.save(path)
    match = Match(
        "docx_paragraph_0",
        EntityType.ORGANIZATION,
        "ООО Тест",
        0,
        8,
        "[ОРГАНИЗАЦИЯ_1]",
        "gigachat",
        0.8,
        Location(paragraph_index=0),
    )
    result = await run_pipeline(
        path,
        tmp_path / "result",
        [EntityType.ORGANIZATION],
        UncertainClient(),
        prepared_matches=[match],
        answers={"party:organization:[ОРГАНИЗАЦИЯ_1]": "buyer"},
    )
    assert result.status is JobStatus.DONE
    assert result.matches[0].party_role is PartyRole.BUYER


@pytest.mark.asyncio
async def test_restart_marks_active_stage_failed(tmp_path):
    db, repo = await repository(tmp_path)
    try:
        await repo.update_stage("abc", "extract", "done")
        await repo.update_stage("abc", "detect", "active")
        assert await repo.recover_interrupted() == 1
        job = await repo.get_job("abc")
        assert job.status is JobStatus.FAILED
        assert job.stages[1]["status"] == "error"
        assert job.progress == 20
    finally:
        await db.close()
