"""Extractor tests for the three required customer formats."""

from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook

from core.extractors.docx_extractor import extract_docx
from core.extractors.ocr.tesseract_provider import TesseractOcrProvider
from core.extractors.pdf_extractor import extract_pdf
from core.extractors.xlsx_extractor import extract_xlsx
from core.models import BlockKind, DocumentFormat


def test_docx_preserves_leading_offsets_and_table_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("  Телефон: +7 999 123-45-67")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "ИНН: 7707083893"
    cell.add_paragraph("email@example.test")
    document.save(path)

    extracted = extract_docx(path)

    assert extracted.format is DocumentFormat.DOCX
    assert extracted.blocks[0].text.startswith("  Телефон")
    table_blocks = [b for b in extracted.blocks if b.kind is BlockKind.DOCX_TABLE_CELL]
    assert [b.location.cell_paragraph_index for b in table_blocks] == [0, 1]


def test_xlsx_preserves_formulas_and_whitespace_offsets(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "  email@example.test  "
    sheet["B1"] = "=1+1"
    workbook.save(path)
    workbook.close()

    extracted = extract_xlsx(path)

    assert extracted.format is DocumentFormat.XLSX
    assert [block.text for block in extracted.blocks] == ["  email@example.test  "]


def test_pdf_marks_image_only_document_as_scanned(tmp_path: Path) -> None:
    text_path = tmp_path / "text.pdf"
    scanned_path = tmp_path / "scan.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "email@example.test")
    document.save(text_path)
    document.close()
    document = fitz.open()
    document.new_page()
    document.save(scanned_path)
    document.close()

    text = extract_pdf(text_path)
    scanned = extract_pdf(scanned_path)

    assert text.format is DocumentFormat.PDF
    assert any("email@example.test" in block.text for block in text.blocks)
    assert text.is_scanned is False
    assert scanned.is_scanned is True
    assert scanned.ocr_pages == (0,)


def test_pdf_marks_only_image_pages_for_ocr(tmp_path: Path) -> None:
    path = tmp_path / "mixed.pdf"
    document = fitz.open()
    document.new_page().insert_text((72, 72), "native text")
    document.new_page()
    document.save(path)
    document.close()

    extracted = extract_pdf(path)

    assert extracted.is_scanned is True
    assert extracted.ocr_pages == (1,)
    assert len(extracted.blocks) == 1


def test_tesseract_tsv_preserves_word_character_coordinates() -> None:
    tsv = (
        "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
        "5\t1\t1\t1\t1\t1\t10\t20\t30\t10\t95\tИванов\n"
        "5\t1\t1\t1\t1\t2\t45\t20\t20\t10\t93\tИ.И.\n"
    )

    blocks = TesseractOcrProvider._parse_tsv(
        tsv,
        2,
        image_width=100,
        image_height=100,
        page_width=200,
        page_height=200,
    )

    assert [block.text for block in blocks] == ["Иванов И.И."]
    assert blocks[0].location.ocr_words == (
        (0, 6, 20.0, 40.0, 80.0, 60.0),
        (7, 11, 90.0, 40.0, 130.0, 60.0),
    )
