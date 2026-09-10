from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook

from core.preview import build_preview


def _all_text(elements: list[dict[str, object]]) -> str:
    values: list[str] = []
    for element in elements:
        if element["type"] == "table":
            for row in element["rows"]:  # type: ignore[index]
                for cell in row:
                    values.extend(segment["text"] for segment in cell["segments"])
        else:
            values.extend(segment["text"] for segment in element["segments"])  # type: ignore[index]
    return "".join(values)


def test_docx_preview_preserves_body_order_and_marks_replacements(tmp_path: Path) -> None:
    path = tmp_path / "result.docx"
    document = Document()
    document.add_paragraph("До [ФИО_1]")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Таблица"
    document.add_paragraph("После")
    document.save(path)

    elements, truncated = build_preview(path, ["[ФИО_1]"])

    assert not truncated
    assert [element["type"] for element in elements] == ["field", "table", "field"]
    assert elements[0]["segments"][1]["replacement"] == "[ФИО_1]"  # type: ignore[index]


def test_xlsx_and_pdf_preview_are_readable_and_support_truncation(tmp_path: Path) -> None:
    xlsx = tmp_path / "result.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "[ИНН_1]"
    workbook.save(xlsx)
    workbook.close()
    xlsx_elements, _ = build_preview(xlsx, ["[ИНН_1]"])
    assert "[ИНН_1]" in _all_text(xlsx_elements)

    pdf = tmp_path / "result.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "preview text")
    document.save(pdf)
    document.close()
    pdf_elements, truncated = build_preview(pdf, [], limit=5)
    assert truncated
    assert len(_all_text(pdf_elements)) <= 5
