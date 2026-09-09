from pathlib import Path

from openpyxl import load_workbook

from core.models import EntityType, Location, Match, PartyRole
from core.report.report_generator import REPORT_COLUMNS, generate_report


def _match(applied: bool, location: Location) -> Match:
    return Match(
        block_id="block",
        entity_type=EntityType.EMAIL,
        text="email@example.test",
        start=0,
        end=18,
        replacement="[EMAIL_1]",
        source="rule",
        confidence=0.97,
        location=location,
        party_role=PartyRole.UNKNOWN,
        applied=applied,
    )


def test_xlsx_report_contains_only_applied_matches(tmp_path: Path) -> None:
    path = tmp_path / "report.xlsx"
    applied = _match(True, Location(sheet_name="Лист1", cell_coordinate="B2"))
    skipped = _match(False, Location(paragraph_index=2))

    assert generate_report({"a": [applied], "b": [skipped]}, path) == path

    workbook = load_workbook(path)
    sheet = workbook.active
    assert tuple(cell.value for cell in sheet[1]) == REPORT_COLUMNS
    assert sheet.max_row == 2
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:I2"
    assert sheet["G2"].number_format == "0.00%"
    assert "Лист1" in sheet["H2"].value
    workbook.close()


def test_empty_xlsx_report_is_openable(tmp_path: Path) -> None:
    path = generate_report({}, tmp_path / "empty.xlsx")
    workbook = load_workbook(path)
    assert workbook.active.max_row == 1
    workbook.close()
